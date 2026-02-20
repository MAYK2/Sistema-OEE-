from flask import Blueprint, render_template, request, redirect, url_for, jsonify, flash
from flask_login import login_user, logout_user, login_required, current_user
from datetime import datetime, timedelta, date
from app.extensions import db
from app.models.work_order import WorkOrder, WorkOrderStatus
from app.models.product import Product
from app.models.client import Client
from app.models.downtime import DowntimeEvent
from app.models.user import User
from app.models.line import Line

views_bp = Blueprint('views', __name__)

from functools import wraps

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            flash("⛔ Acceso denegado: Se requieren permisos de Administrador.", "danger")
            return redirect(url_for('views.index'))
        return f(*args, **kwargs)
    return decorated_function

@views_bp.route('/')
def index():
    if not current_user.is_authenticated:
        return redirect(url_for('views.create_order'))
    # Solo buscamos órdenes que el operario pueda ver para trabajar
    # PENDING: Creadas pero no iniciadas
    # EXECUTION: Órdenes que ya están corriendo (cronómetro activo)
    orders = WorkOrder.query.filter(
        WorkOrder.status.in_([WorkOrderStatus.PENDING, WorkOrderStatus.EXECUTION])
    ).all()
    
    return render_template('index.html', orders=orders)

@views_bp.route('/tutorial')
@login_required
def tutorial():
    return render_template('tutorial.html')

@views_bp.route('/orders/new', methods=['GET', 'POST'])
def create_order():
    if request.method == 'POST':
        if request.is_json:
            data = request.get_json()
            
            # --- Lógica de Cliente (Tu código original) ---
            client_mode = data.get('client_mode')
            client_id = None
            
            if client_mode == 'existing':
                client_id = data.get('existing_client_id')
            elif client_mode == 'new':
                new_client_name = data.get('new_client_name')
                new_client_phone = data.get('new_client_phone')
                new_client_email = data.get('new_client_email')
                new_client = Client(name=new_client_name, phone=new_client_phone, email=new_client_email)
                db.session.add(new_client)
                db.session.commit()
                client_id = new_client.id
            
            # --- Fecha Estimada ---
            estimated_finish_str = data.get('estimated_finish')
            estimated_finish = None
            if estimated_finish_str:
                try:
                    estimated_finish = datetime.strptime(estimated_finish_str, '%Y-%m-%dT%H:%M')
                except ValueError:
                    pass

            # --- Crear Órdenes ---
            items = data.get('items', [])
            created_count = 0
            
            # Optimization: Fetch last OT number ONCE for the batch
            year = datetime.now().year
            last_order = WorkOrder.query.filter(WorkOrder.ot_number.like(f'OT-{year}-%'))\
                                        .order_by(WorkOrder.id.desc())\
                                        .first()
            
            last_seq = 0
            if last_order:
                try:
                    last_seq = int(last_order.ot_number.split('-')[-1])
                except ValueError:
                    pass

            for item in items:
                # Increment sequence for each item in the batch
                new_seq = last_seq + 1 + created_count
                ot_number = f"OT-{year}-{new_seq:04d}"
                
                planned_quantity = item.get('quantity')
                
                new_order = WorkOrder(
                    ot_number=ot_number,
                    client_id=client_id,
                    product_id=item.get('product_id'),
                    planned_quantity=planned_quantity,
                    estimated_finish=estimated_finish,
                    status=WorkOrderStatus.PENDING
                )
                db.session.add(new_order)
                created_count += 1
            
            db.session.commit()
            return jsonify({'message': f'{created_count} órdenes creadas correctamente'}), 200
        
        return redirect(url_for('views.index'))

    # GET request
    products = Product.query.all()
    clients = Client.query.all()
    today_date = datetime.now().strftime('%Y-%m-%dT%H:%M')
    return render_template('create_order.html', products=products, clients=clients, today_date=today_date)

@views_bp.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('views.index'))
    
    error = None
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = User.query.filter_by(username=username).first()
        
        if user and user.check_password(password):
            login_user(user)
            return redirect(url_for('views.index'))
        else:
            error = "Usuario o contraseña incorrectos"
            
    return render_template('login.html', error=error)

@views_bp.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('views.login'))

@views_bp.route('/operate/<int:order_id>')
@login_required
def operate(order_id):
    # FORCE FRESH SESSION to ensure we see the latest downtimes/status changes
    db.session.remove()
    
    order = WorkOrder.query.get_or_404(order_id)
    
    # Calcular Tiempos de Parada
    downtimes = DowntimeEvent.query.filter_by(work_order_id=order.id).all()
    total_downtime_seconds = sum(d.duration_seconds for d in downtimes if d.duration_seconds)
    
    # Verificar si hay una pausa activa ahora mismo (Prioritize latest if multiple exist)
    open_downtimes = [d for d in downtimes if d.end_time is None]
    open_downtimes.sort(key=lambda x: x.start_time, reverse=True)
    current_pause = open_downtimes[0] if open_downtimes else None
    
    pause_start_iso = None
    is_changeover = False
    
    if current_pause and current_pause.start_time:
        pause_start_iso = current_pause.start_time.isoformat()
        if current_pause.reason == "Cambio de Paso":
            is_changeover = True

    # CORRECCIÓN: Usar 'start_time' estándar en lugar de '_real'
    # Calculate elapsed seconds server-side for robust stopwatch
    initial_elapsed_seconds = 0
    now_time = datetime.now()
    
    # SELF-HEALING: If order is EXECUTION but has no start_time (Data Corruption Fix)
    if order.status == WorkOrderStatus.EXECUTION and not order.start_time:
         from datetime import timedelta
         # Recovery: Set start time to now so timer resumes/starts
         order.start_time = now_time
         db.session.commit()

    if order.start_time:
         # If order is finished, clamp to end_time
         if order.status == WorkOrderStatus.FINISHED and order.end_time:
             now_time = order.end_time
             
         # If we are currently paused, the clock shouldn't be "ticking" past the pause start
         # But the total_downtime calculation already sums CLOSED downtimes.
         # So: Elapsed = (Now - Start) - ClosedDowntimes - CurrentOpenDowntimeDuration
         
         total_duration = (now_time - order.start_time).total_seconds()
         
         current_pause_duration = 0
         if current_pause and current_pause.start_time:
              current_pause_duration = (now_time - current_pause.start_time).total_seconds()
              
         initial_elapsed_seconds = int(max(0, total_duration - total_downtime_seconds - current_pause_duration))
         
         initial_elapsed_seconds = int(max(0, total_duration - total_downtime_seconds - current_pause_duration))
         
    # Log unconditionally to debug why start_time is missing
    try:
         with open('/tmp/matrix_debug.log', 'a') as f:
             f.write(f"DEBUG_VIEW: OT={order.ot_number} ID={order.id}\n")
             f.write(f"   Status={order.status.name} StartTime={order.start_time}\n")
             f.write(f"   Now={now_time} Elapsed={initial_elapsed_seconds}\n")
    except Exception as e:
         pass

    return render_template('operator.html', 
                           order=order, 
                           initial_elapsed_seconds=initial_elapsed_seconds,
                           pause_start_iso=pause_start_iso,
                           is_changeover=is_changeover)

# --- API ENDPOINTS ---

@views_bp.route('/api/orders/<int:id>/start', methods=['POST'])
@login_required
def start_order(id):
    # FORCE FRESH SESSION to avoid stale reads (Isolation Level issues)
    # Ensure any previous transaction state is cleared so we see updates from other requests (like verify/pause)
    db.session.remove() 

    order = WorkOrder.query.get_or_404(id)
    data = request.get_json() or {}
    
    # ROBUSTNESS: Always check for open downtimes and close them when starting
    # This ensures that even if status transition logic was bypassed/buggy, we don't leave open pauses ticking
    # Get the LATEST open downtime (in case of multiple zombies)
    open_downtime = DowntimeEvent.query.filter_by(work_order_id=order.id, end_time=None)\
                                       .order_by(DowntimeEvent.start_time.desc())\
                                       .first()
    
    if open_downtime:
        open_downtime.end_time = datetime.now()
        # Calculate duration
        delta = open_downtime.end_time - open_downtime.start_time
        open_downtime.duration_seconds = int(delta.total_seconds())
        # Commit immediately to ensure consistency
        db.session.add(open_downtime)
        db.session.commit()

    # Logic to start order if PENDING (New or Paused)
    if order.status == WorkOrderStatus.PENDING:
        # Change status to EXECUTION
        order.status = WorkOrderStatus.EXECUTION
        
        # Record start time only if it's the first start
        if not order.start_time: 
            order.start_time = datetime.now()
            
        # Save operators count if provided
        if 'operators_count' in data:
            order.operators_count = int(data['operators_count'])
            
        db.session.commit()
        
        with open('/tmp/matrix_debug.log', 'a') as f:
             f.write(f"DEBUG_START: OT={order.ot_number} ID={order.id}\n")
             f.write(f"   Set StartTime={order.start_time}\n")
        
        return jsonify({
            "message": "Order started", 
            "status": order.status.value, 
            "start_time": order.start_time.isoformat()
        })
    
    # Si la orden ya estaba en ejecución, no hacemos nada
    return jsonify({"error": "Invalid status transition"}), 400

@views_bp.route('/api/orders/<int:id>/pause', methods=['POST'])
@login_required
def pause_order(id):
    order = WorkOrder.query.get_or_404(id)
    data = request.get_json()
    reason = data.get('reason', 'Sin motivo especificado')

    if order.status == WorkOrderStatus.EXECUTION or order.status == WorkOrderStatus.IN_PROGRESS:
        order.status = WorkOrderStatus.PENDING # Usamos PENDING como estado de Pausa visual
        
        # Crear evento de parada
        downtime = DowntimeEvent(
            work_order_id=order.id,
            reason=reason,
            comment=data.get('comment'),
            start_time=datetime.now()
        )
        db.session.add(downtime)
        db.session.commit()
        return jsonify({"message": "Order paused", "status": order.status.value})
        
    return jsonify({"error": "Order must be in EXECUTION to pause"}), 400

@views_bp.route('/api/orders/<int:id>/finish', methods=['POST'])
@login_required
def finish_order(id):
    order = WorkOrder.query.get_or_404(id)
    data = request.get_json() or {}
    
    # Guardar cantidad final ingresada por el usuario
    if 'produced_quantity' in data:
         order.total_produced = int(data['produced_quantity'])

    order.status = WorkOrderStatus.FINISHED
    # CORRECCIÓN: Usar 'end_time' estándar
    order.end_time = datetime.now()
    
    # Cerrar pausas abiertas si quedaron
    downtime = DowntimeEvent.query.filter_by(work_order_id=order.id, end_time=None).first()
    if downtime:
        downtime.end_time = order.end_time
        downtime.duration_seconds = int((downtime.end_time - downtime.start_time).total_seconds())

    db.session.commit()
    return jsonify({"message": "Order finished", "status": order.status.value})

# --- CHANGEOVER ENDPOINTS ---
@views_bp.route('/api/orders/<int:id>/changeover/start', methods=['POST'])
@login_required
def start_changeover(id):
    # Force Session Refresh
    db.session.remove()

    order = WorkOrder.query.get_or_404(id)
    
    # Check if already in changeover or running
    if order.status != WorkOrderStatus.PENDING:
         return jsonify({"error": "Order must be PENDING to start changeover"}), 400

    # Check for open events
    open_event = DowntimeEvent.query.filter_by(work_order_id=order.id, end_time=None).first()
    if open_event:
        return jsonify({"error": "An event is already in progress"}), 400

    # Create Changeover Event
    event = DowntimeEvent(
        work_order_id=order.id,
        reason="Cambio de Paso",
        start_time=datetime.now()
    )
    db.session.add(event)
    db.session.commit()
    
    return jsonify({"message": "Changeover started", "start_time": event.start_time.isoformat()})

@views_bp.route('/api/orders/<int:id>/changeover/stop', methods=['POST'])
@login_required
def stop_changeover(id):
    # Force Session Refresh
    db.session.remove()

    order = WorkOrder.query.get_or_404(id)
    
    # Find the open changeover event
    open_event = DowntimeEvent.query.filter_by(work_order_id=order.id, end_time=None).first()
    
    if not open_event:
        return jsonify({"message": "No active changeover found"}), 200
        
    open_event.end_time = datetime.now()
    delta = open_event.end_time - open_event.start_time
    open_event.duration_seconds = int(delta.total_seconds())
    
    db.session.commit()
    
    return jsonify({"message": "Changeover stopped", "duration": open_event.duration_seconds})

# --- REPORTES ---

@views_bp.route('/config/lines')
@login_required
@admin_required
def config_lines():
    return render_template('config_lines.html')

@views_bp.route('/history')
@login_required
@admin_required
def history():
    orders = WorkOrder.query.filter_by(status=WorkOrderStatus.FINISHED)\
                            .order_by(WorkOrder.end_time.desc())\
                            .all()
    return render_template('history.html', orders=orders)

@views_bp.route('/report/<int:id>')
@login_required
@admin_required
def report_detail(id):
    order = WorkOrder.query.get_or_404(id)
    
    # Generate Chronological Activity Log
    timeline = []
    
    # Start point
    current_pivot = order.start_time
    
    # Sort events just in case
    events = sorted(order.downtime_events, key=lambda x: x.start_time)
    
    order_end = order.end_time if order.end_time else datetime.now()

    for dt in events:
        # 1. Identify Production Block before this downtime
        # If there is a gap between current_pivot and dt.start_time, that's Production
        if dt.start_time > current_pivot:
            prod_duration = (dt.start_time - current_pivot).total_seconds()
            if prod_duration > 1: # Ignore micro-gaps
                timeline.append({
                    'type': 'production',
                    'start_time': current_pivot,
                    'end_time': dt.start_time,
                    'reason': 'Producción (En Ejecución)',
                    'duration_seconds': int(prod_duration),
                    'color': 'success',
                    'icon': 'bi-gear-fill spin-slow'
                })
        
        # 2. Add the Downtime Block
        dt_end = dt.end_time if dt.end_time else order_end
        timeline.append({
            'type': 'downtime',
            'start_time': dt.start_time,
            'end_time': dt.end_time, # Can be None if active
            'reason': dt.reason,
            'comment': dt.comment,
            'duration_seconds': dt.duration_seconds if dt.end_time else int((datetime.now() - dt.start_time).total_seconds()),
            'color': 'info' if dt.reason == 'Cambio de Paso' else 'danger',
            'icon': 'bi-tools' if dt.reason == 'Cambio de Paso' else 'bi-pause-circle-fill'
        })
        
        # Move pivot
        current_pivot = dt_end
    
    # 3. Check for final Production Block (after last downtime)
    if current_pivot and current_pivot < order_end:
         prod_duration = (order_end - current_pivot).total_seconds()
         if prod_duration > 1:
             timeline.append({
                    'type': 'production',
                    'start_time': current_pivot,
                    'end_time': order.end_time, # None if running
                    'reason': 'Producción (En Ejecución)',
                    'duration_seconds': int(prod_duration),
                    'color': 'success',
                    'icon': 'bi-gear-fill spin-slow'
            })
            
    return render_template('report_detail.html', order=order, timeline=timeline)

@views_bp.route('/reports')
@login_required
@admin_required
def reports():
    # Filtros
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    client_id = request.args.get('client_id')

    query = WorkOrder.query.filter_by(status=WorkOrderStatus.FINISHED)

    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
            query = query.filter(WorkOrder.end_time >= start_date)
        except ValueError:
            pass
    
    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d')
            end_date = end_date.replace(hour=23, minute=59, second=59)
            query = query.filter(WorkOrder.end_time <= end_date)
        except ValueError:
            pass

    if client_id:
        query = query.filter_by(client_id=int(client_id))

    orders = query.all()
    clients = Client.query.all()

    # Lógica de Agregación para Reportes
    final_report = get_report_data_helper(orders)
    
    return render_template('reports.html', report_data=final_report, clients=clients)

def get_report_data_helper(orders):
    report_data = {} 
    for order in orders:
        pid = order.product_id
        if pid not in report_data:
            report_data[pid] = {
                'product_name': order.product.name,
                'orders_count': 0,
                'total_produced': 0,
                'total_run_time': 0,
                'total_downtime': 0
            }
        
        # Usamos propiedades seguras del modelo o cálculo manual si falla
        run_time = getattr(order, 'run_time_seconds', 0)
        dt_seconds = sum(d.duration_seconds for d in order.downtime_events if d.duration_seconds)
        
        report_data[pid]['orders_count'] += 1
        report_data[pid]['total_produced'] += order.total_produced
        report_data[pid]['total_run_time'] += run_time
        report_data[pid]['total_downtime'] += dt_seconds

    final_report = []
    for pid, data in report_data.items():
        final_report.append({
            'product_name': data['product_name'],
            'orders_count': data['orders_count'],
            'total_produced': data['total_produced'],
            'run_time_minutes': round(data['total_run_time'] / 60, 1),
            'downtime_minutes': round(data['total_downtime'] / 60, 1)
        })
    return final_report

@views_bp.route('/reports/export')
@login_required
@admin_required
def export_reports():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    from flask import send_file

    # --- Reusing Filter Logic (Duplicated slightly for standalone nature, or extract to helper if used 3+ times) ---
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    client_id = request.args.get('client_id')

    query = WorkOrder.query.filter_by(status=WorkOrderStatus.FINISHED)

    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
            query = query.filter(WorkOrder.end_time >= start_date)
        except ValueError:
            pass
    
    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d')
            end_date = end_date.replace(hour=23, minute=59, second=59)
            query = query.filter(WorkOrder.end_time <= end_date)
        except ValueError:
            pass

    if client_id:
        query = query.filter_by(client_id=int(client_id))

    orders = query.all()
    data = get_report_data_helper(orders)

    # --- Create Excel ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Matrix OEE"

    # Add Period Info Row
    period_text = "Periodo: Historial Completo"
    if start_date_str and end_date_str:
        period_text = f"Periodo: {start_date_str} al {end_date_str}"
    elif start_date_str:
        period_text = f"Periodo: Desde {start_date_str}"
    elif end_date_str:
        period_text = f"Periodo: Hasta {end_date_str}"

    ws.append([period_text])
    ws.merge_cells('A1:E1') # Merge across 5 columns
    
    title_cell = ws['A1']
    title_cell.font = Font(size=12, bold=True, italic=True)
    title_cell.alignment = Alignment(horizontal="center")

    # Headers (Now on Row 2)
    headers = ["Producto", "N° Órdenes", "Total Producido", "Tiempo Ejecución (min)", "Tiempo Paradas (min)"]
    ws.append(headers)

    # Style Header
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid") # Dark Blue
    
    for cell in ws[2]: # Row 2 is headers now
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data
    for row in data:
        ws.append([
            row['product_name'],
            row['orders_count'],
            row['total_produced'],
            row['run_time_minutes'],
            row['downtime_minutes']
        ])

    # Auto-adjust columns
    for col in ws.columns:
        max_length = 0
        # Use Row 2 (Header) to get column letter to avoid MergedCell in Row 1
        # col[0] is Row 1, col[1] is Row 2
        try:
            column = get_column_letter(col[1].column)
        except:
             column = get_column_letter(col[0].column)

        for cell in col:
            # Skip Row 1 (Title) so it doesn't mess up widths for columns A-E
            if cell.row == 1:
                continue
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    # Generate Stream
    excel_stream = BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    
    filename = f"Reporte_Produccion_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return send_file(
        excel_stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

from app.services.time_sync import get_network_time

@views_bp.route('/dashboard')
@login_required
@admin_required
def dashboard():
    # Use Network Time for accurate Dashboard Date/Day
    now_network = get_network_time()
    today = now_network.date()
    
    start_day = datetime.combine(today, datetime.min.time())
    start_week = start_day - timedelta(days=today.weekday())
    start_month = start_day.replace(day=1)
    
    # 1. Fetch All Lines
    lines = Line.query.all()
    if not lines:
        pass 

    # --- HELPER FUNCTIONS (Refactored for Line Context) ---

    def get_stats(start_date, line_id):
        orders = WorkOrder.query.join(Product).filter(
            WorkOrder.status == WorkOrderStatus.FINISHED,
            WorkOrder.end_time >= start_date,
            Product.line_id == line_id
        ).all()
        
        stats = {}
        for order in orders:
            p_name = order.product.name
            if p_name not in stats:
                stats[p_name] = 0
            stats[p_name] += order.total_produced
        return stats

    def get_downtime_stats(start_date, line_id):
        downtimes = DowntimeEvent.query.join(WorkOrder).join(Product).filter(
            DowntimeEvent.start_time >= start_date,
            DowntimeEvent.reason != None,
            Product.line_id == line_id
        ).all()

        reason_counts = {}
        total_events = 0

        for dt in downtimes:
            r = dt.reason
            if r not in reason_counts:
                reason_counts[r] = 0
            reason_counts[r] += 1
            total_events += 1
            
        reason_stats = {}
        if total_events > 0:
            for r, count in reason_counts.items():
                reason_stats[r] = round((count / total_events) * 100, 1)
        
        return reason_stats

    def get_time_stats(start_date, line_id):
        orders = WorkOrder.query.join(Product).filter(
            WorkOrder.status == WorkOrderStatus.FINISHED,
            WorkOrder.end_time >= start_date,
            Product.line_id == line_id
        ).all()
        
        sum_run_seconds = 0
        sum_changeover_seconds = 0
        sum_stops_seconds = 0
        
        for order in orders:
            sum_run_seconds += getattr(order, 'run_time_seconds', 0)
            for d in order.downtime_events:
                if d.duration_seconds:
                    if d.reason == "Cambio de Paso":
                        sum_changeover_seconds += d.duration_seconds
                    else:
                        sum_stops_seconds += d.duration_seconds
            
        return {
            'T. Ejecución': round(sum_run_seconds / 60, 1),
            'T. Cambio de Paso': round(sum_changeover_seconds / 60, 1),
            'T. Paradas': round(sum_stops_seconds / 60, 1)
        }

    def get_timeline_data(target_date, line_id):
        next_day = target_date + timedelta(days=1)
        orders = WorkOrder.query.join(Product).filter(
            WorkOrder.start_time < next_day,
            (WorkOrder.end_time >= target_date) | (WorkOrder.end_time == None),
            Product.line_id == line_id
        ).all()
        
        timeline_events = []
        
        for order in orders:
            effective_start = max(order.start_time, target_date)
            effective_end = order.end_time if order.end_time else datetime.now()
            effective_end = min(effective_end, next_day)
            
            if effective_end <= effective_start:
                continue

            timeline_events.append({
                'x': [effective_start.isoformat(), effective_end.isoformat()],
                'y': 'Producción', 
                'fillColor': '#27ae60', 
                'label': f"{order.product.name} (OT-{order.ot_number.split('-')[-1] if order.ot_number else order.id})",
                'type': 'Production'
            })
            
            for d in order.downtime_events:
                d_start = max(d.start_time, target_date)
                d_end = d.end_time if d.end_time else datetime.now()
                d_end = min(d_end, next_day)
                
                if d_end <= d_start:
                    continue
                
                fill_color = '#95a5a6'
                r_lower = d.reason.lower() if d.reason else "otro"
                if "cambio" in r_lower: fill_color = '#2980b9'
                elif "falla" in r_lower: fill_color = '#c0392b'
                elif "insumo" in r_lower: fill_color = '#e67e22'
                elif "limpieza" in r_lower: fill_color = '#f1c40f'
                elif "otro" in r_lower: fill_color = '#9b59b6'
                
                timeline_events.append({
                    'x': [d_start.isoformat(), d_end.isoformat()],
                    'y': 'Producción',
                    'fillColor': fill_color,
                    'label': d.reason,
                    'type': 'Downtime'
                })
        
        return timeline_events

    def get_downtime_pareto(start_date, line_id):
        orders = WorkOrder.query.join(Product).filter(
            WorkOrder.end_time >= start_date,
            Product.line_id == line_id
        ).all()
        reasons = {}
        
        for order in orders:
            for d in order.downtime_events:
                if d.duration_seconds and d.reason != "Cambio de Paso":
                    if d.reason not in reasons:
                        reasons[d.reason] = 0
                    reasons[d.reason] += d.duration_seconds
        
        sorted_reasons = sorted(reasons.items(), key=lambda item: item[1], reverse=True)
        total_downtime = sum(reasons.values())
        
        return {
            'labels': [r[0] for r in sorted_reasons],
            'data': [round(r[1] / 60, 1) for r in sorted_reasons],
            'total_minutes': round(total_downtime / 60, 1)
        }


    def calculate_availability(start_date, end_date, shift_start_time, shift_end_time, line_id):
        """
        Calculates availability based on overlapping production intervals vs planned shift time.
        """
        # 1. Fetch Orders in Period
        orders = WorkOrder.query.join(Product).filter(
            WorkOrder.status == WorkOrderStatus.FINISHED,
            WorkOrder.start_time < end_date, 
            WorkOrder.end_time >= start_date,
            Product.line_id == line_id
        ).all()
        
        # Also include active orders? calculating until 'now' or 'end_date'
        active_orders = WorkOrder.query.join(Product).filter(
            WorkOrder.status == WorkOrderStatus.EXECUTION,
            WorkOrder.start_time < end_date,
            Product.line_id == line_id
        ).all()
        
        all_orders = orders + active_orders
        
        # 2. Extract Intervals (Clamped to Period)
        intervals = []
        for o in all_orders:
            # Handle potentially missing times
            s = o.start_time
            e = o.end_time if o.end_time else datetime.now()
            
            # Clamp to query window
            eff_s = max(s, start_date)
            eff_e = min(e, end_date)
            
            if eff_e > eff_s:
                intervals.append((eff_s, eff_e))
        
        # 3. Merge Intervals
        if not intervals:
            return 0
            
        intervals.sort(key=lambda x: x[0])
        
        merged = []
        if intervals:
            curr_start, curr_end = intervals[0]
            for next_start, next_end in intervals[1:]:
                if next_start < curr_end: # Overlap
                    curr_end = max(curr_end, next_end)
                else:
                    merged.append((curr_start, curr_end))
                    curr_start, curr_end = next_start, next_end
            merged.append((curr_start, curr_end))
            
        total_run_time_seconds = sum((end - start).total_seconds() for start, end in merged)
        
        # 4. Calculate Planned Time (Shift Hours * Days in Period)
        # Simple approximation: Count business days? Or just Days passed?
        # User wants "Disponibilidad Proredio".
        # Let's count "Active Days" (days with at least 1 order) to be fair?
        # Or standard calendar days?
        # If we check "Month", strict shift hours (8h * 30d) might be too much if they don't work weekends.
        # Let's use the period duration but clamped to Shift Hours.
        
        # Iterate through days in range
        total_planned_seconds = 0
        current_day = start_date.date()
        end_day_date = end_date.date()
        
        while current_day <= end_day_date:
            # Shift Start/End for this day
            s_dt = datetime.combine(current_day, shift_start_time)
            e_dt = datetime.combine(current_day, shift_end_time)
            
            # Clamp to global start/end (e.g. if start_date is mid-day)
            # Actually dashboard standard start_date is 00:00.
            
            # Check if now is before shift end (for Today)
            # If historical, full shift.
            
            # Optimization: Check if this day is in the past or today
            day_start_limit = max(s_dt, start_date)
            day_end_limit = min(e_dt, end_date)
            
            if day_end_limit > day_start_limit:
                 total_planned_seconds += (day_end_limit - day_start_limit).total_seconds()
            
            current_day += timedelta(days=1)
            
        if total_planned_seconds == 0:
            return 0
            
        return round(min(100, (total_run_time_seconds / total_planned_seconds) * 100), 1)

    # --- MAIN LOOP PER LINE ---
    dashboard_data = []

    for line in lines:
        line_data = {
            'id': line.id,
            'name': line.name,
            'shift': {
                'start': line.shift_start.strftime('%H:%M') if line.shift_start else "08:00",
                'end': line.shift_end.strftime('%H:%M') if line.shift_end else "17:00"
            }
        }

        # 1. Stats per Period
        line_data['day'] = get_stats(start_day, line.id)
        line_data['week'] = get_stats(start_week, line.id)
        line_data['month'] = get_stats(start_month, line.id)
        
        line_data['downtime_month'] = get_downtime_stats(start_month, line.id)
        line_data['time_week'] = get_time_stats(start_week, line.id)
        line_data['time_month'] = get_time_stats(start_month, line.id)
        line_data['pareto'] = get_downtime_pareto(start_month, line.id)
        line_data['timeline'] = get_timeline_data(start_day, line.id)

        # 2. OEE / Availability Calculation
        shift_start_time = line.shift_start if line.shift_start else datetime.time(8,0)
        shift_end_time = line.shift_end if line.shift_end else datetime.time(17,0)
        
        # Today
        line_data['oee_percent'] = calculate_availability(start_day, datetime.now(), shift_start_time, shift_end_time, line.id)
        
        # Week (From start of week to NOW)
        line_data['oee_week'] = calculate_availability(start_week, datetime.now(), shift_start_time, shift_end_time, line.id)
        
        # Month (From start of month to NOW)
        line_data['oee_month'] = calculate_availability(start_month, datetime.now(), shift_start_time, shift_end_time, line.id)
        
        dashboard_data.append(line_data)

    date_names = {
        'day': today.strftime("%d/%m"),
        'month': today.strftime("%B")
    }

    target_data = {
        'target': 1000,
        'current': sum([sum(l['month'].values()) for l in dashboard_data]),
        'percentage': 0,
        'status': 'En Camino'
    }
    
    return render_template('dashboard.html', 
                           lines_data=dashboard_data, 
                           target_data=target_data,
                           date_names=date_names,
                           stats={})

@views_bp.route('/general-analytics')
@login_required
def general_analytics():
    # --- ANALÍTICA DE PRODUCCIÓN (GLOBAL ESTIMADA) ---
    window_days = 30
    last_30_days_date = datetime.now() - timedelta(days=window_days)
    
    orders_30d = WorkOrder.query.filter(
        WorkOrder.status == WorkOrderStatus.FINISHED,
        WorkOrder.end_time >= last_30_days_date
    ).all()
    
    total_products = {} 
    total_liters_30d = 0
    category_stats = {'sodas': 0, 'bidones_10': 0, 'bidones_20': 0}

    for o in orders_30d:
        if o.product:
            p_name = o.product.name
            qty = o.total_produced or 0
            name_lower = p_name.lower()
            volume = 0
            
            if "soda" in name_lower or "sifon" in name_lower: 
                volume = 1
                category_stats['sodas'] += qty
            elif "bidon" in name_lower or "botellon" in name_lower:
                if "10" in name_lower:
                    volume = 10
                    category_stats['bidones_10'] += qty
                else:
                    volume = 20
                    category_stats['bidones_20'] += qty
            
            total_liters_30d += (qty * volume)

            if p_name not in total_products:
                total_products[p_name] = 0
            total_products[p_name] += qty

    return render_template('general_analytics.html', 
                           total_liters_30d=int(total_liters_30d), 
                           category_stats=category_stats)
