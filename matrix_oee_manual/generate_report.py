from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_report():
    doc = Document()
    
    # --- Styles ---
    # Title
    title = doc.add_heading('INFORME DE GESTIÓN TÉCNICO Y FUNCIONAL', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Proyecto: Sistema Matrix OEE - Optimización de Procesos y Digitalización de Datos')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(14)
    
    doc.add_paragraph() # Spacer

    # --- Content ---
    
    # 1. Executive Summary
    doc.add_heading('1. Resumen Ejecutivo: ¿Qué estamos solucionando?', level=1)
    p = doc.add_paragraph('El objetivo de Matrix OEE es eliminar la incertidumbre en la planta de producción. Transformamos la gestión actual, basada en registros manuales y estimaciones, en un ecosistema de digitalización completa en tiempo real.')
    
    p = doc.add_paragraph('El sistema centraliza la operación conectando la Planificación (Gerencia) con la Ejecución (Operarios), permitiendo detectar fugas de tiempo, desperdicios y cuellos de botella al instante. No solo medimos cuánto producimos, sino cómo lo producimos.')

    # 2. Workflow
    doc.add_heading('2. Flujo de Trabajo y Estandarización', level=1)
    doc.add_paragraph('El sistema estandariza el proceso productivo en tres etapas críticas, asegurando la trazabilidad del dato desde la oficina hasta la planta.')

    # 2.1 Planning
    doc.add_heading('2.1. Etapa de Planificación (Interfaz Administrativa)', level=2)
    
    # Image 1
    try:
        doc.add_picture('app/static/img/tutorial/create_order.png', width=Inches(6))
        last_paragraph = doc.paragraphs[-1] 
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = doc.add_paragraph('Imagen 1: Interfaz de Nueva Orden')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(9)
    except Exception as e:
        doc.add_paragraph(f"[Error al cargar imagen: {e}]")

    doc.add_paragraph('Acción: Es el punto de partida. El usuario administrativo carga la Orden de Trabajo (OT).')
    doc.add_paragraph('Datos Clave: Se define el cliente, el producto, la cantidad objetivo y la línea asignada.')
    doc.add_paragraph('Resultado: La máquina queda "programada" digitalmente. Se elimina el papel y las instrucciones verbales confusas.')

    # 2.2 Operational
    doc.add_heading('2.2. Etapa Operativa (Interfaz de Planta / Tablet-First)', level=2)

    # Image 2 (Using Dashboard Summary or generic placeholder if "List" isn't perfect, but text says "Listado". 
    # I will use order_history.png as it looks like a list, and clarify in caption if needed, or stick to user mapping
    # User's Image 2 was likely intended to be the Operator View. I will use 'login.png' or 'dashboard' if I lack the operator specific one.
    # Wait, I have 'dashboard_summary.png'. 
    # Let's use 'order_history.png' as "Listado" for now as it's the best list I have.
    try:
        doc.add_picture('app/static/img/tutorial/order_history.png', width=Inches(6))
        last_paragraph = doc.paragraphs[-1] 
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = doc.add_paragraph('Imagen 2: Listado de Órdenes')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(9)
    except:
        pass

    doc.add_paragraph('Acceso: El operario se loguea en la línea con una interfaz optimizada para tablets (botones grandes, sin scroll, alta usabilidad).')
    
    doc.add_heading('Funciones Críticas:', level=3)
    p = doc.add_paragraph()
    p.add_run('Inicio de Producción: ').bold = True
    p.add_run('Activación de cronómetros de tiempo productivo con un solo toque.')
    
    p = doc.add_paragraph()
    p.add_run('Registro de Paradas: ').bold = True
    p.add_run('Bloqueo de "paradas fantasma". El sistema obliga a seleccionar un motivo (Mecánico, Eléctrico, Insumos) al pausar.')
    
    p = doc.add_paragraph()
    p.add_run('Justificación: ').bold = True
    p.add_run('Si se selecciona "Otros", se exige justificación escrita obligatoria.')
    
    p = doc.add_paragraph()
    p.add_run('Cambio de Paso: ').bold = True
    p.add_run('Función dedicada para medir tiempos de configuración (Setup) entre productos.')
    
    p = doc.add_paragraph()
    p.add_run('Persistencia: ').bold = True
    p.add_run('Protección total de datos ante recargas o micro-cortes de red.')

    # 2.3 Analysis
    doc.add_heading('2.3. Etapa de Análisis (Tableros y Reportes)', level=2)
    doc.add_paragraph('Transformación de datos crudos en decisiones visuales.')
    
    p = doc.add_paragraph()
    p.add_run('Timeline en Vivo: ').bold = True
    p.add_run('Visualización minuto a minuto del estado de la máquina (Verde: Produciendo, Rojo: Parada, Azul: Setup).')
    
    p = doc.add_paragraph()
    p.add_run('Gráficos de Pareto: ').bold = True
    p.add_run('Identificación automática del 20% de las causas que generan el 80% de las pérdidas.')
    
    p = doc.add_paragraph()
    p.add_run('Consumo de Recursos: ').bold = True
    p.add_run('Estimación de consumo de agua e insumos críticos por OT.')

    # 3. KPIs
    doc.add_heading('3. Indicadores de Desempeño (KPIs)', level=1)
    doc.add_paragraph('El software automatiza el cálculo del estándar de oro industrial, el OEE (Eficiencia General de los Equipos):')
    
    p = doc.add_paragraph()
    p.add_run('Disponibilidad: ').bold = True
    p.add_run(' (Tiempo productivo / Tiempo disponible). ¿La máquina funcionó o estuvo parada?')
    
    p = doc.add_paragraph()
    p.add_run('Rendimiento: ').bold = True
    p.add_run(' (Velocidad real / Velocidad teórica). ¿Fuimos rápidos o lentos según el estándar de la máquina?')
    
    p = doc.add_paragraph()
    p.add_run('Calidad: ').bold = True
    p.add_run(' (Piezas buenas / Total producido). ¿Cuánto desperdicio generamos?')

    # Image 3
    try:
        doc.add_picture('app/static/img/tutorial/config_lines.png', width=Inches(6))
        last_paragraph = doc.paragraphs[-1] 
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = doc.add_paragraph('Imagen 3: Configuración de Línea y Turnos')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(9)
    except:
        pass

    doc.add_paragraph('Nota: Se incluye el indicador de "Cumplimiento del Plan", una proyección en tiempo real sobre si se alcanzará la meta del turno actual.')

    # 4. Technical Improvements
    doc.add_heading('4. Mejoras Técnicas y Estabilidad', level=1)
    doc.add_paragraph('¿Por qué confiar en la arquitectura de Matrix?')
    
    p = doc.add_paragraph()
    p.add_run('Cálculos Reales: ').bold = True
    p.add_run('Algoritmos basados en tiempo transcurrido real, eliminando distorsiones y promedios falsos al inicio de las órdenes.')
    
    p = doc.add_paragraph()
    p.add_run('Integridad de Sesión: ').bold = True
    p.add_run('Corrección de errores de "reseteo"; la base de datos mantiene la identidad del operario y el estado de la máquina aunque se cierre el navegador.')
    
    p = doc.add_paragraph()
    p.add_run('UX/UI Ágil: ').bold = True
    p.add_run('Sistema de navegación por pestañas (Vivo, Tendencias, Análisis) para acceso inmediato a la información sin tiempos de carga excesivos.')

    # Image 4
    try:
        doc.add_picture('app/static/img/tutorial/login.png', width=Inches(4))
        last_paragraph = doc.paragraphs[-1] 
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = doc.add_paragraph('Imagen 4: Login y Seguridad')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(9)
    except:
        pass

    # 5. Conclusion
    doc.add_heading('5. Conclusión y Valor Agregado', level=1)
    doc.add_paragraph('El Sistema Matrix OEE evoluciona el rol del registro de producción: deja de ser una tarea burocrática para convertirse en un asistente de productividad.')
    doc.add_paragraph('Al estandarizar la entrada de datos, eliminamos la subjetividad humana y obtenemos una "radiografía" exacta de la fábrica. Esto permite a la gerencia dejar de "apagar incendios" basados en suposiciones y comenzar a tomar decisiones estratégicas basadas en datos reales.')

    # Image 5
    try:
        # Reusing order history or dashboard summary as appropriate. 
        # Text says "Historial de Órdenes".
        doc.add_picture('app/static/img/tutorial/order_history.png', width=Inches(6))
        last_paragraph = doc.paragraphs[-1] 
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = doc.add_paragraph('Imagen 5: Historial de Órdenes')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(9)
    except:
        pass

    # Save
    doc.save('Informe_Gestion_Matrix_OEE.docx')
    print("Documento guardado como Informe_Gestion_Matrix_OEE.docx")

if __name__ == "__main__":
    create_report()
